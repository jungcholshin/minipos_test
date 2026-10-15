"""Импорт каталога из Excel / CSV."""
import io

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker

from app import config, importer, services
from app.db import Base
from app.models import Category, Product, StockMovement


@pytest.fixture
def db(monkeypatch):
    monkeypatch.setattr(config, "RETAIL_MARKUP", 1.5)
    monkeypatch.setattr(config, "RETAIL_ROUND", 100)
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    s = sessionmaker(bind=engine, expire_on_commit=False)()
    c = Category(name="SPF")
    s.add(c)
    s.flush()
    s.add(Product(name="Beauty of Joseon Relief Sun", category_id=c.id, barcode="8809738316139",
                  purchase_price=3000, wholesale_price=4000, selling_price=6000, stock_quantity=5))
    s.commit()
    yield s
    s.close()


def xlsx(rows, extra_sheet_first=False):
    wb = Workbook()
    ws = wb.active
    if extra_sheet_first:
        ws.append(["ID", "Дата", "Сумма"])
        ws = wb.create_sheet("Products")
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_template_is_readable():
    wb = load_workbook(io.BytesIO(importer.template()))
    assert [c.value for c in wb["Товары"][1]][:3] == ["Название", "Категория", "Штрихкод"]


def test_preview_create_update_error(db):
    data = xlsx([
        ["Наименование", "Категория", "Штрих-код", "Закупка", "Оптовая цена", "Розничная цена", "Кол-во"],
        ["Новый крем", "Кремы", 8801234567890, "2 500", 3000, None, 10],     # опт → розница 4500
        ["Тонер", None, None, None, None, "7 000 ₸", 3],                       # только розница
        ["Beauty of joseon relief sun", None, None, None, 4330, None, 20],     # обновление по названию
        [None, "x", None, None, 100, None, None],                              # нет названия
        ["Без цены", "x", None, None, None, None, None],                       # нет цены
    ])
    out = importer.import_products(db, data, "c.xlsx", do_apply=False)
    assert out["summary"] == {"create": 2, "update": 1, "same": 0, "error": 2}
    items = out["items"]
    assert items[0]["selling_price"] == 4500 and items[0]["barcode"] == "8801234567890"
    assert items[1]["wholesale_price"] == 4670 and items[1]["category"] == "Без категории"
    assert db.execute(select(Product)).scalars().all().__len__() == 1      # предпросмотр ничего не меняет


def test_apply(db):
    data = xlsx([["Название", "Категория", "Опт", "Остаток"],
                 ["Новый крем", "Кремы", 3000, 10],
                 ["Beauty of Joseon Relief Sun", "", 4330, 20]])
    out = importer.import_products(db, data, "c.xlsx", do_apply=True)
    assert out["summary"]["create"] == 1 and out["summary"]["update"] == 1
    new = db.execute(select(Product).where(Product.name == "Новый крем")).scalar_one()
    assert new.selling_price == 4500 and new.stock_quantity == 10 and new.category.name == "Кремы"
    old = db.get(Product, 1)
    assert old.wholesale_price == 4330 and old.stock_quantity == 20 and old.purchase_price == 3000
    moves = db.execute(select(StockMovement)).scalars().all()
    assert {m.movement_type for m in moves} == {"PURCHASE", "ADJUSTMENT"}


def test_match_by_barcode_and_export_sheet(db):
    data = xlsx([["ID", "Штрихкод", "Название", "Категория", "Закупка", "Опт", "Розница", "Остаток"],
                 [1, "8809738316139", "Другое имя", "SPF", 3000, 4000, 6000, 5]], extra_sheet_first=True)
    out = importer.import_products(db, data, "export.xlsx", do_apply=False)
    assert out["items"][0]["action"] == "update"   # найден по штрихкоду, лист Products выбран сам


def test_csv_semicolon(db):
    data = "Название;Розница;Остаток\nМаска;2 400;30\n".encode("utf-8-sig")
    out = importer.import_products(db, data, "c.csv", do_apply=True)
    assert out["summary"]["create"] == 1
    assert db.execute(select(Product).where(Product.name == "Маска")).scalar_one().wholesale_price == 1600


def test_bad_file(db):
    with pytest.raises(services.BusinessError):
        importer.import_products(db, xlsx([["Цена", "Кол-во"], [1, 2]]), "c.xlsx", do_apply=False)
    with pytest.raises(services.BusinessError):
        importer.import_products(db, b"xx", "c.pdf", do_apply=False)
